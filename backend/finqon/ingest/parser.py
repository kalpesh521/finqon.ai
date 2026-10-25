# ruff: noqa: E501
"""Deterministic statement parsing: CSV, TSV, XLSX and text-layer PDF tables.

Pure functions with no database or network access. The LLM never extracts amounts;
unknown layouts are rejected with a clear message instead of guessed at. When the file
has a running balance column the rows are reconciled against it, and a file whose
balances do not add up is quarantined by the caller.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import pairwise
from typing import Any

import pdfplumber
from openpyxl import load_workbook

SUPPORTED_SUFFIXES = (".csv", ".tsv", ".txt", ".xlsx", ".pdf")
MAX_ROWS = 20_000

_DATE_FORMATS = (
    "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%y", "%d %b %Y", "%d-%b-%Y",
    "%d-%b-%y", "%d %B %Y", "%b %d, %Y", "%Y/%m/%d", "%d.%m.%Y",
)  # fmt: skip

_HEADERS: dict[str, tuple[str, ...]] = {
    "date": ("date", "txn date", "transaction date", "value date", "posting date", "tran date"),
    "description": ("description", "narration", "particulars", "details", "remarks",
                    "transaction remarks", "transaction details", "merchant"),
    "debit": ("debit", "withdrawal", "withdrawals", "withdrawal amt", "dr", "debit amount", "paid out"),
    "credit": ("credit", "deposit", "deposits", "deposit amt", "cr", "credit amount", "paid in"),
    "amount": ("amount", "transaction amount", "amt"),
    "type": ("type", "dr/cr", "cr/dr", "direction", "transaction type"),
    "balance": ("balance", "closing balance", "running balance", "available balance"),
}  # fmt: skip


class ParseError(ValueError):
    """The file cannot be read as a bank statement. The message is shown to the user."""


@dataclass(frozen=True, slots=True)
class ParsedRow:
    txn_date: date
    description: str
    amount: Decimal  # always positive
    direction: str  # "debit" | "credit"
    balance: Decimal | None


@dataclass(slots=True)
class ParsedStatement:
    rows: list[ParsedRow]
    skipped: int = 0
    warnings: list[str] = field(default_factory=list)
    reconciled: bool | None = None  # None when the file has no balance column


def parse_statement(filename: str, data: bytes) -> ParsedStatement:
    name = filename.lower()
    if name.endswith(".xlsx"):
        grid = _read_xlsx(data)
    elif name.endswith(".pdf"):
        grid = _read_pdf(data)
    elif name.endswith((".csv", ".tsv", ".txt")):
        grid = _read_delimited(data)
    else:
        raise ParseError("Unsupported file type. Upload CSV, XLSX or PDF.")
    return rows_from_grid(grid)


# ---------------------------------------------------------------------------- readers


def _read_delimited(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise ParseError("The file is not readable text.")
    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("The Excel file could not be opened.") from exc
    sheet = book.active
    if sheet is None:
        raise ParseError("The Excel file has no sheets.")
    return [
        [_cell_text(c) for c in row]
        for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 50)
    ]


def _read_pdf(data: bytes) -> list[list[str]]:
    grid: list[list[str]] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:200]:
                for table in page.extract_tables():
                    grid.extend([_cell_text(c) for c in row] for row in table)
    except Exception as exc:
        raise ParseError("The PDF could not be read.") from exc
    if not grid:
        raise ParseError(
            "No transaction table found in this PDF. Scanned PDFs are not supported yet; "
            "export a CSV or a text-based PDF from your bank."
        )
    return grid


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).replace("\n", " ").strip()


# --------------------------------------------------------------------------- normalise


def rows_from_grid(grid: list[list[str]]) -> ParsedStatement:
    header_index, columns = _find_header(grid)
    out = ParsedStatement(rows=[])
    for raw in grid[header_index + 1 :]:
        if not any(raw):
            continue
        if len(out.rows) >= MAX_ROWS:
            out.warnings.append(f"Only the first {MAX_ROWS:,} rows were imported.")
            break
        row = _to_row(raw, columns)
        if row is None:
            out.skipped += 1
        else:
            out.rows.append(row)
    if not out.rows:
        raise ParseError("No valid transactions were found in this file.")
    if out.skipped:
        out.warnings.append(f"{out.skipped} row(s) were skipped (blank, totals or unreadable).")
    out.reconciled = _reconcile(out.rows)
    if out.reconciled is False:
        out.warnings.append("Running balances do not match the transaction amounts.")
    return out


def _norm(value: str) -> str:
    return re.sub(r"[^a-z/ ]", "", value.lower()).strip()


def _find_header(grid: list[list[str]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(grid[:40]):
        columns: dict[str, int] = {}
        for position, cell in enumerate(row):
            label = _norm(cell)
            if not label:
                continue
            for key, names in _HEADERS.items():
                if key not in columns and label in names:
                    columns[key] = position
        has_money = "amount" in columns or "debit" in columns or "credit" in columns
        if "date" in columns and "description" in columns and has_money:
            return index, columns
    raise ParseError(
        "Could not find the transaction columns. The file needs a date, a description and "
        "an amount (or debit and credit) column."
    )


def parse_amount(value: str) -> Decimal | None:
    cleaned = re.sub(r"[₹$€£,\s]|rs\.?|inr", "", value.strip(), flags=re.IGNORECASE)
    if not cleaned or cleaned in {"-", "--"}:
        return None
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    cleaned = cleaned.strip("()")
    suffix = None
    if cleaned.lower().endswith(("dr", "cr")):
        suffix = cleaned[-2:].lower()
        cleaned = cleaned[:-2]
    if not re.fullmatch(r"-?\d+(\.\d+)?", cleaned):
        return None
    try:
        number = Decimal(cleaned)
    except InvalidOperation:
        return None
    if not number.is_finite():
        return None
    if negative or suffix == "dr":
        number = -abs(number)
    return number


def parse_date(value: str) -> date | None:
    text = value.strip()
    if not text:
        return None
    text = text.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", text) else text
    for fmt in _DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        if 2000 <= parsed.year <= date.today().year + 1:
            return parsed
    return None


def _to_row(raw: list[str], columns: dict[str, int]) -> ParsedRow | None:
    def get(key: str) -> str:
        position = columns.get(key)
        return raw[position] if position is not None and position < len(raw) else ""

    txn_date = parse_date(get("date"))
    description = " ".join(get("description").split())
    if txn_date is None or not description:
        return None

    amount: Decimal | None
    direction: str
    if "debit" in columns or "credit" in columns:
        debit = parse_amount(get("debit")) if "debit" in columns else None
        credit = parse_amount(get("credit")) if "credit" in columns else None
        if debit and abs(debit) > 0:
            amount, direction = abs(debit), "debit"
        elif credit and abs(credit) > 0:
            amount, direction = abs(credit), "credit"
        else:
            return None
    else:
        signed = parse_amount(get("amount"))
        if signed is None or signed == 0:
            return None
        marker = get("type").lower()
        if marker.startswith(("dr", "debit", "d")) and marker:
            direction = "debit"
        elif marker.startswith(("cr", "credit", "c")) and marker:
            direction = "credit"
        else:
            direction = "debit" if signed < 0 else "credit"
        amount = abs(signed)

    if amount >= Decimal("1e12"):
        return None
    balance = parse_amount(get("balance")) if "balance" in columns else None
    return ParsedRow(
        txn_date, description[:500], amount.quantize(Decimal("0.01")), direction, balance
    )


def _reconcile(rows: list[ParsedRow], tolerance: Decimal = Decimal("0.02")) -> bool | None:
    """Check that ``balance[i] - balance[i-1]`` equals the signed amount.

    Statements are either oldest-first or newest-first, so both orders are tried and the
    better one wins. Passing needs at least 95% of the comparable pairs to agree.
    """
    with_balance = [r for r in rows if r.balance is not None]
    if len(with_balance) < max(3, len(rows) // 2):
        return None

    def agreement(ordered: list[ParsedRow]) -> float:
        good = total = 0
        for prev, cur in pairwise(ordered):
            if prev.balance is None or cur.balance is None:
                continue
            signed = cur.amount if cur.direction == "credit" else -cur.amount
            total += 1
            if abs((cur.balance - prev.balance) - signed) <= tolerance:
                good += 1
        return good / total if total else 0.0

    best = max(agreement(rows), agreement(list(reversed(rows))))
    return best >= 0.95
